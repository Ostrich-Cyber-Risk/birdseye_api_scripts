import argparse
import re
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from getpass import getpass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import requests
from openpyxl import load_workbook

from ostrichApi import OstrichApi, Assessment, BusinessUnit

PRODUCTION_BASE_URL = 'https://api.ostrichcyber-risk.com'
DEFAULT_SHEET_NAME = 'CIS-Internal Controls'
DEFAULT_ASSESSMENT_TYPE_ID = 'CIS 8.1'
DEFAULT_COVERAGE = 80

SAFEGUARD_COLUMN = 'CIS Safeguard'
STATUS_COLUMN = 'Assessment Status'
BAND_COLUMN = 'Safeguard Score'
OBSERVATIONS_COLUMN = 'Observations'

SCORE = 'score'
UNKNOWN = 'unknown'
SKIPPED = 'skipped'

# The workbook records an implementation judgement per safeguard, not a percentage. These are
# the five values its Instructions sheet allows, and the process score each one becomes.
# A safeguard the workbook scores is imported at that score; "Not Applicable" carries no score
# and is left untouched rather than being marked ignored, because ignoring applies to every
# assessment of the type on the business unit, not just this one.
STATUS_CONVERSION: Dict[str, Tuple[str, Optional[int]]] = {
    'met': (SCORE, 100),
    'partially met': (SCORE, 75),
    'not met': (SCORE, 0),
    'not assessed': (SCORE, 0),
    'not applicable': (SKIPPED, None),
}

# The workbook also carries a 1-5 band derived from the status. It is not the source of truth
# (band 1 means both "Not Met" and "Not Assessed"), but a row whose band disagrees with its
# status has been hand-edited, which is worth surfacing before an upload.
EXPECTED_BAND: Dict[str, Optional[int]] = {
    'met': 5,
    'partially met': 4,
    'not met': 1,
    'not assessed': 1,
    'not applicable': None,
}

SAFEGUARD_PATTERN = re.compile(r'Safeguard\s+(\d+\.\d+)\b')
BAND_PATTERN = re.compile(r'^([1-5])\b')

# The api validates assessment dates against exactly this layout, all nine fractional digits:
# emu's AddAssessmentRequestAssessment carries
# `validate:"datetime=2006-01-02T15:04:05.000000000Z"`. A plain RFC-3339 stamp is rejected as
# "Invalid Request Data format" before the handler runs.
NANOSECOND_TIME_FORMAT = '%Y-%m-%dT%H:%M:%S.%f000Z'


class Row:
    """One safeguard read out of the workbook, with the conversion already applied."""

    def __init__(self, number: str, status: str, kind: str, score: Optional[int], observations: str):
        self.number = number
        self.status = status
        self.kind = kind
        self.score = score
        self.observations = observations


def main():
    major_version = sys.version_info.major
    minor_version = sys.version_info.minor
    if major_version != 3 or minor_version < 12:
        raise Exception(f"Running in Python {major_version}.{minor_version} ... Minimum required Python version is 3.12")

    args = parse_args()

    print()
    print('=' * 78)
    print(' Import a CIS internal-controls workbook into a Birdseye assessment')
    print('=' * 78)
    print()
    print('This walks through six steps and asks before it writes anything:')
    print('  1. pick the workbook and sheet        4. pick the business unit')
    print('  2. review the score conversion        5. pick or create the assessment')
    print('  3. choose the coverage value          6. review and upload')
    print()

    step('Step 1 of 6: workbook and sheet')
    workbook_path = args.workbook or prompt_required('Path to the workbook (.xlsx)')
    sheet_name = choose_sheet(workbook_path, args.sheet, args.yes)
    rows = read_workbook(workbook_path, sheet_name)
    review_conversion(rows, args.yes)
    coverage = choose_coverage(args.coverage, args.yes)

    api_client = authenticate(args.apiKey.strip() if args.apiKey else read_api_key(), args.baseUrl)
    if args.baseUrl != PRODUCTION_BASE_URL:
        print(f'Using {args.baseUrl}')

    business_unit_id = choose_business_unit(api_client, args.businessUnitId, args.yes)
    assessment = choose_assessment(api_client, business_unit_id, args.assessmentTypeId,
                                   args.assessmentName, args.yes, args.dryRun)
    if assessment is None:
        print('\nNothing was created. Re-run without --dryRun to create the assessment.')
        return

    print(f"\nUsing assessment {assessment['assessmentId']} ({assessment['assessmentName']})")
    content = api_client.get_assessment_content(business_unit_id, assessment['assessmentId'])
    safeguard_to_question = build_safeguard_map(content)
    print(f'The assessment content exposes {len(safeguard_to_question)} CIS safeguards.')

    unmatched = sorted(r.number for r in rows if r.number not in safeguard_to_question)
    if unmatched:
        raise SystemExit(
            f'\n{len(unmatched)} safeguards in the workbook are not in this assessment '
            f'content: {unmatched}\nCheck the assessment type is the CIS version the workbook '
            f'was built against before importing a partial set.')

    scores = build_scores(rows, safeguard_to_question, coverage)
    notes = [] if args.skipNotes else build_notes(rows, safeguard_to_question)

    step('Step 6 of 6: review before upload')
    summarise_upload(rows, scores, notes, coverage)

    if args.dryRun:
        print('\nDry run, nothing was written. Re-run without --dryRun to upload.')
        return

    if not args.yes:
        confirm_or_exit(f"Upload to \"{assessment['assessmentName']}\" now?")

    upload(api_client, business_unit_id, assessment, scores, notes)
    print('\nDone.')


def step(title: str) -> None:
    print()
    print('-' * 78)
    print(f' {title}')
    print('-' * 78)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description='Import CIS safeguard scores from an internal-controls workbook into a '
                    'Birdseye assessment, one business unit per run. Prompts through every '
                    'step; pass the flags below to answer any of them up front.')
    parser.add_argument('--workbook', help='Path to the .xlsx workbook.')
    parser.add_argument('--sheet', help=f'Sheet holding the safeguards (default: "{DEFAULT_SHEET_NAME}").')
    parser.add_argument('--businessUnitId', help='Id of the (leaf) business unit to upload to.')
    parser.add_argument('--assessmentName', help='Assessment to reuse, or to create if it does not exist.')
    parser.add_argument('--assessmentTypeId', default=DEFAULT_ASSESSMENT_TYPE_ID,
                        help=f'Assessment content type (default: "{DEFAULT_ASSESSMENT_TYPE_ID}").')
    parser.add_argument('--coverage', type=parse_coverage,
                        help='Coverage paired with every process score: a multiple of 5 in 0-100, '
                             f'or "unknown" to leave it unanswered (default: {DEFAULT_COVERAGE}).')
    parser.add_argument('--apiKey', help='Api key. Omit to be prompted without echoing it.')
    parser.add_argument('--baseUrl', default=PRODUCTION_BASE_URL,
                        help=f'Api base url (default: {PRODUCTION_BASE_URL}). Point this at a '
                             f'non-production environment to rehearse an import.')
    parser.add_argument('--skipNotes', action='store_true',
                        help=f'Do not import the "{OBSERVATIONS_COLUMN}" column as notes.')
    parser.add_argument('--dryRun', action='store_true', help='Report what would happen without writing.')
    parser.add_argument('--yes', action='store_true',
                        help='Accept every prompt. Requires the answers to be passed as flags.')
    return parser.parse_args()


def parse_coverage(value: str):
    """Parse a coverage argument into a multiple-of-5 int in 0-100, or UNKNOWN."""
    text = value.strip().lower()
    if text in {UNKNOWN, 'idk'}:
        return UNKNOWN
    try:
        number = int(text)
    except ValueError:
        raise argparse.ArgumentTypeError(
            f'coverage must be an integer 0-100 in steps of 5, or "unknown" (got {value!r})')
    if number < 0 or number > 100 or number % 5 != 0:
        raise argparse.ArgumentTypeError(
            f'coverage must be a multiple of 5 between 0 and 100, or "unknown" (got {value!r})')
    return number


def authenticate(api_key: str, base_url: str) -> OstrichApi:
    """Exchange the api key for a token, reporting a rejected key as a message not a traceback."""
    try:
        return OstrichApi(api_key=api_key, base_url=base_url)
    except requests.HTTPError as error:
        status = error.response.status_code if error.response is not None else None
        if status in (401, 403):
            raise SystemExit('That Api Key was rejected. Keys expire, so check it is still '
                             'current in Birdseye under the business unit Api Keys.')
        raise


def read_api_key() -> str:
    """Read the api key without echoing it, falling back to stdin when there is no console.

    getpass reads the console directly on Windows, so it would hang forever on a piped or
    redirected stdin rather than reading the key that was piped in.
    """
    prompt = '\nEnter your Api Key (not echoed): '
    if sys.stdin.isatty():
        return getpass(prompt).strip()
    print(prompt.strip())
    return sys.stdin.readline().strip()


def prompt_required(label: str) -> str:
    while True:
        value = input(f'{label}: ').strip().strip('"')
        if value:
            return value
        print('  Required.')


def confirm_or_exit(question: str) -> None:
    """Ask a yes-or-no question that defaults to yes, and exit unless the answer is yes."""
    answer = input(f'\n{question} [Y/n] ').strip().lower()
    if answer not in {'', 'y', 'yes'}:
        raise SystemExit('Stopped, nothing was written.')


def choose_sheet(workbook_path: str, sheet: Optional[str], assume_yes: bool) -> str:
    sheet_names = load_workbook(workbook_path, read_only=True).sheetnames
    if sheet:
        if sheet not in sheet_names:
            raise SystemExit(f'No sheet named "{sheet}". This workbook has: {sheet_names}')
        print(f'Sheet: {sheet}')
        return sheet
    if assume_yes:
        raise SystemExit('--yes needs --sheet.')

    default = DEFAULT_SHEET_NAME if DEFAULT_SHEET_NAME in sheet_names else None
    print(f'\n{Path(workbook_path).name} has {len(sheet_names)} sheets:')
    for index, name in enumerate(sheet_names, start=1):
        marker = '  <- default' if name == default else ''
        print(f'  {index:2d}. {name}{marker}')
    while True:
        suffix = ' [Enter for the default]' if default else ''
        answer = input(f'Which sheet holds the safeguards?{suffix} ').strip()
        if not answer and default:
            return default
        if answer.isdigit() and 1 <= int(answer) <= len(sheet_names):
            return sheet_names[int(answer) - 1]
        if answer in sheet_names:
            return answer
        print('  Enter a number from the list.')


def read_workbook(workbook_path: str, sheet_name: str) -> List[Row]:
    """Read one safeguard per row, converting each Assessment Status into a process score.

    Raises on an unrecognised status rather than skipping the row: a status the workbook's own
    Instructions sheet does not allow means the file is not what this importer expects, and a
    silently dropped safeguard is worse than a stopped import.
    """
    sheet = load_workbook(workbook_path, data_only=True)[sheet_name]
    grid = list(sheet.iter_rows(values_only=True))
    if not grid:
        raise SystemExit(f'Sheet "{sheet_name}" is empty.')

    header = {str(value).strip(): index for index, value in enumerate(grid[0]) if value is not None}
    for required in (SAFEGUARD_COLUMN, STATUS_COLUMN):
        if required not in header:
            raise SystemExit(f'Sheet "{sheet_name}" has no "{required}" column. Found: {sorted(header)}')

    def cell(values, column: str) -> str:
        index = header.get(column)
        if index is None or index >= len(values) or values[index] is None:
            return ''
        return str(values[index]).strip()

    rows: List[Row] = []
    unrecognised: List[str] = []
    band_mismatches: List[str] = []
    for values in grid[1:]:
        number = cell(values, SAFEGUARD_COLUMN)
        status = cell(values, STATUS_COLUMN)
        if not number:
            continue
        conversion = STATUS_CONVERSION.get(status.lower())
        if conversion is None:
            unrecognised.append(f'{number}: {status or "(blank)"}')
            continue
        kind, score = conversion
        if band_disagrees(cell(values, BAND_COLUMN), status):
            band_mismatches.append(f'{number}: status "{status}", {BAND_COLUMN} "{cell(values, BAND_COLUMN)}"')
        rows.append(Row(number, status, kind, score, cell(values, OBSERVATIONS_COLUMN)))

    if unrecognised:
        raise SystemExit(
            f'{len(unrecognised)} rows have a status outside {sorted(STATUS_CONVERSION)}:\n  '
            + '\n  '.join(unrecognised))
    if not rows:
        raise SystemExit(f'Sheet "{sheet_name}" has no safeguard rows.')

    counts = Counter(r.number for r in rows)
    duplicates = sorted(number for number, count in counts.items() if count > 1)
    if duplicates:
        raise SystemExit(f'Sheet "{sheet_name}" lists these safeguards more than once: {duplicates}')

    print(f'\nRead {len(rows)} safeguards from "{sheet_name}".')
    if band_mismatches:
        print(f'WARNING: {len(band_mismatches)} rows where {BAND_COLUMN} does not match '
              f'{STATUS_COLUMN}. {STATUS_COLUMN} is what gets imported:', file=sys.stderr)
        for line in band_mismatches:
            print(f'  {line}', file=sys.stderr)
    return rows


def band_disagrees(band_cell: str, status: str) -> bool:
    """True when the workbook's 1-5 band contradicts the status it is supposed to be derived from."""
    if not band_cell:
        return False
    match = BAND_PATTERN.match(band_cell)
    if match is None:
        return False
    return EXPECTED_BAND.get(status.lower()) not in (None, int(match.group(1)))


def review_conversion(rows: List[Row], assume_yes: bool) -> None:
    step('Step 2 of 6: how each status becomes a score')
    print('Birdseye scores the "process" aspect of each safeguard from 0 to 100 in steps of 25.')
    print(f'The workbook records a judgement per safeguard, so "{STATUS_COLUMN}" is converted:')
    print()
    print(f'  {STATUS_COLUMN:<16} {"becomes":<28} {"in this workbook":>16}')
    for status, (kind, score) in STATUS_CONVERSION.items():
        label = {
            SCORE: f'process score {score}',
            UNKNOWN: 'unanswered ("I don\'t know")',
            SKIPPED: 'left untouched, no score',
        }[kind]
        count = sum(1 for r in rows if r.status.lower() == status)
        print(f'  {status.title():<16} {label:<28} {count:>16}')
    print()
    print('A safeguard the workbook scores is imported at that score. "Not Applicable" carries')
    print('no score in the workbook, so it is left alone rather than marked ignored: ignoring')
    print('applies to every assessment of this type on the business unit, not just this one.')
    if not assume_yes:
        confirm_or_exit('Is that the conversion you want?')


def choose_coverage(coverage, assume_yes: bool):
    step('Step 3 of 6: coverage')
    if coverage is not None:
        print(f'Coverage: {coverage}')
        return coverage
    if assume_yes:
        return DEFAULT_COVERAGE

    print('Every safeguard has a second "coverage" aspect: how much of the organisation the')
    print('control actually reaches. It scales the process score down, so 100 applies no')
    print('reduction and a lower value discounts the score proportionally.')
    print('The workbook does not measure coverage, so one value is applied to every safeguard.')
    print(f'Enter a multiple of 5 in 0-100, or "unknown" to leave it unanswered (excluded from')
    print(f'the average). Press Enter for {DEFAULT_COVERAGE}.')
    while True:
        answer = input('Coverage: ').strip()
        if not answer:
            return DEFAULT_COVERAGE
        try:
            return parse_coverage(answer)
        except argparse.ArgumentTypeError as error:
            print(f'  {error}')


def choose_business_unit(api_client: OstrichApi, business_unit_id: Optional[str], assume_yes: bool) -> str:
    step('Step 4 of 6: business unit')
    if business_unit_id:
        print(f'Business unit: {business_unit_id}')
        return business_unit_id
    if assume_yes:
        raise SystemExit('--yes needs --businessUnitId.')

    print('Retrieving the business units your Api Key can see...')
    business_units = flatten_business_units(api_client.get_business_units())
    if not business_units:
        raise SystemExit('This Api Key cannot see any business units.')
    return pick_business_unit(business_units)['businessUnitId']


def pick_business_unit(business_units: List[BusinessUnit]) -> BusinessUnit:
    """Offer the business units the key can see, the first one being the default."""
    for index, business_unit in enumerate(business_units, start=1):
        marker = '  <- default' if index == 1 else ''
        print(f"  {index:2d}. {business_unit.get('name', '')}  "
              f"({business_unit.get('businessUnitId', '')}){marker}")
    print('Scores need the PractitionerAssessments role on the leaf business unit, and creating')
    print('an assessment needs Manager. Roles do not inherit from a parent.')
    while True:
        answer = input('Which business unit? [Enter for the default] ').strip()
        if not answer:
            return business_units[0]
        if answer.isdigit() and 1 <= int(answer) <= len(business_units):
            return business_units[int(answer) - 1]
        print(f'  Enter a number from 1 to {len(business_units)}.')


def choose_assessment(api_client: OstrichApi, business_unit_id: str, assessment_type_id: str,
                      assessment_name: Optional[str], assume_yes: bool,
                      dry_run: bool) -> Optional[Assessment]:
    step('Step 5 of 6: assessment')
    existing = [a for a in api_client.get_assessments(business_unit_id)
                if a.get('assessmentTypeId') == assessment_type_id]

    if assessment_name is None:
        if assume_yes:
            raise SystemExit('--yes needs --assessmentName.')
        chosen = pick_assessment(existing, assessment_type_id)
        if chosen is not None:
            return update_existing(chosen, assume_yes)
        assessment_name = prompt_required('Name for the new assessment')

    matches = [a for a in existing if a.get('assessmentName') == assessment_name]
    if len(matches) > 1:
        raise SystemExit(f'{len(matches)} assessments named "{assessment_name}" of type '
                         f'"{assessment_type_id}" on this business unit; names must be unique.')
    if matches:
        return update_existing(matches[0], assume_yes)

    print(f'\nWould create a new "{assessment_type_id}" assessment named "{assessment_name}".')
    if dry_run:
        return None
    if not assume_yes:
        confirm_or_exit('Create it?')
    start = datetime.now(timezone.utc)
    api_client.create_assessment(business_unit_id, assessment_name, assessment_type_id,
                                 api_timestamp(start), api_timestamp(start + timedelta(days=365)))
    # Create returns a bare success message with no id, so re-list to find what it made.
    created = [a for a in api_client.get_assessments(business_unit_id)
               if a.get('assessmentName') == assessment_name
               and a.get('assessmentTypeId') == assessment_type_id]
    if not created:
        raise SystemExit(f'Assessment "{assessment_name}" was created but could not be found when re-listing.')
    return created[0]


def pick_assessment(existing: List[Assessment], assessment_type_id: str) -> Optional[Assessment]:
    """Offer the assessments already on the business unit, or None to create a new one.

    The first existing assessment is the default, so a business unit with just one takes a
    single Enter.
    """
    if not existing:
        print(f'This business unit has no "{assessment_type_id}" assessments yet.')
        return None

    plural = '' if len(existing) == 1 else 's'
    print(f'This business unit has {len(existing)} "{assessment_type_id}" assessment{plural}:')
    for index, assessment in enumerate(existing, start=1):
        marker = '  <- default' if index == 1 else ''
        print(f"  {index:2d}. {assessment.get('assessmentName')}  "
              f"({assessment.get('assessmentId')}){marker}")
    create_option = len(existing) + 1
    print(f'  {create_option:2d}. Create a new assessment')
    print('Updating an existing assessment overwrites the scores of every safeguard in the '
          'workbook.')

    while True:
        answer = input('Which assessment? [Enter for the default] ').strip()
        if not answer:
            return existing[0]
        if answer.isdigit():
            choice = int(answer)
            if choice == create_option:
                return None
            if 1 <= choice <= len(existing):
                return existing[choice - 1]
        print(f'  Enter a number from 1 to {create_option}.')


def update_existing(assessment: Assessment, assume_yes: bool) -> Assessment:
    """Confirm an overwrite of an assessment that already holds scores."""
    print(f"\nReusing existing assessment {assessment['assessmentId']} "
          f"({assessment.get('assessmentName')}).")
    if not assume_yes:
        typed = input('This overwrites its existing scores. Type UPDATE to continue: ').strip()
        if typed != 'UPDATE':
            raise SystemExit('Stopped, nothing was written.')
    return assessment


def api_timestamp(moment: datetime) -> str:
    """Format a datetime the way the assessment api demands. See NANOSECOND_TIME_FORMAT."""
    return moment.strftime(NANOSECOND_TIME_FORMAT)


def build_safeguard_map(content: dict) -> Dict[str, str]:
    """Map a CIS safeguard number (e.g. "1.1") to its Birdseye question id.

    The safeguard number cannot be derived from the question id (question ids are ordered by
    NIST CSF function), so it is parsed from each question's text.
    """
    nodes = content.get('nodes') or {}
    mapping: Dict[str, str] = {}
    for question_id, node in nodes.items():
        if not isinstance(node, dict) or 'process' not in node:
            continue
        match = SAFEGUARD_PATTERN.search(str(node.get('text', '')))
        if match:
            mapping[match.group(1)] = question_id
    return mapping


def build_scores(rows: List[Row], safeguard_to_question: Dict[str, str], coverage) -> List[dict]:
    """Build the score payload. Every process score is paired with a coverage score.

    A safeguard the workbook does not score is left out entirely, so an existing answer in the
    assessment is not overwritten by a guess.
    """
    scores: List[dict] = []
    for row in rows:
        if row.kind == SKIPPED:
            continue
        question_id = safeguard_to_question[row.number]
        process_aspect_id = f'{question_id}-PROCESS'
        if row.kind == UNKNOWN:
            scores.append({'aspectId': process_aspect_id, 'unknown': True})
        else:
            scores.append({'aspectId': process_aspect_id, 'score': row.score})
        scores.append(coverage_score(question_id, coverage))
    return scores


def coverage_score(question_id: str, coverage) -> dict:
    """Build the coverage entry paired with a safeguard's process score.

    Sends either "score" or "unknown", never both: the API rejects an entry carrying both.
    """
    aspect_id = f'{question_id}-COVERAGE'
    if coverage == UNKNOWN:
        return {'aspectId': aspect_id, 'unknown': True}
    return {'aspectId': aspect_id, 'score': coverage}


def build_notes(rows: List[Row], safeguard_to_question: Dict[str, str]) -> List[dict]:
    return [{'aspectId': f'{safeguard_to_question[row.number]}-PROCESS', 'value': row.observations}
            for row in rows if row.observations]


def summarise_upload(rows: List[Row], scores: List[dict], notes: List[dict], coverage) -> None:
    scored = [r for r in rows if r.kind == SCORE]
    print(f'  {len(scored):>4} safeguards scored')
    for value in (100, 75, 50, 25, 0):
        count = sum(1 for r in scored if r.score == value)
        if count:
            print(f'       {count:>4} at {value}')
    skipped = [r for r in rows if r.kind == SKIPPED]
    if skipped:
        print(f'  {len(skipped):>4} left untouched, the workbook gives them no score: '
              f'{", ".join(r.number for r in skipped)}')
    print(f'  {len(scored):>4} coverage aspects set to '
          f'{"unanswered" if coverage == UNKNOWN else coverage}')
    print(f'  {len(notes):>4} notes from "{OBSERVATIONS_COLUMN}"')
    print(f'\n  {len(scores)} score entries in total. First few:')
    for score in scores[:4]:
        value = 'unanswered' if score.get('unknown') else score['score']
        print(f"    {score['aspectId']:<24} {value}")
    if notes:
        print(f"  First note: {notes[0]['aspectId']} {notes[0]['value'][:60]!r}...")


def upload(api_client: OstrichApi, business_unit_id: str, assessment: Assessment,
           scores: List[dict], notes: List[dict]) -> None:
    assessment_id = assessment['assessmentId']
    if notes:
        api_client.save_notes(business_unit_id, assessment_id, notes)
        print(f'Saved {len(notes)} notes.')
    if scores:
        api_client.save_scores(business_unit_id, assessment_id, scores)
        print(f'Uploaded {len(scores)} score entries.')


def flatten_business_units(root_business_units: List[BusinessUnit]) -> List[BusinessUnit]:
    business_units: List[BusinessUnit] = root_business_units.copy()
    for business_unit in root_business_units:
        business_units.extend(flatten_business_units(business_unit.get('businessUnits', [])))
    return business_units


if __name__ == '__main__':
    main()
