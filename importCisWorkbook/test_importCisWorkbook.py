import argparse
import unittest
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook

from importCisWorkbook import (SCORE, SKIPPED, UNKNOWN, Row, api_timestamp, band_disagrees,
                               build_notes, build_safeguard_map, build_scores, confirm_or_exit,
                               coverage_score, parse_coverage, pick_assessment,
                               pick_business_unit, read_workbook)

HEADER = ['CIS Safeguard', 'Control', 'Assessment Status', 'Safeguard Score', 'Observations']


def write_workbook(directory: str, rows, sheet_name: str = 'CIS-Internal Controls') -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    path = str(Path(directory) / 'workbook.xlsx')
    workbook.save(path)
    return path


class ParseCoverageTest(unittest.TestCase):
    def test_accepts_multiples_of_five(self):
        self.assertEqual(80, parse_coverage('80'))
        self.assertEqual(0, parse_coverage('0'))
        self.assertEqual(100, parse_coverage('100'))

    def test_accepts_unknown_tokens(self):
        self.assertEqual(UNKNOWN, parse_coverage('unknown'))
        self.assertEqual(UNKNOWN, parse_coverage('IDK'))

    def test_rejects_off_scale_and_non_numeric(self):
        for value in ('81', '-5', '105', 'high', ''):
            with self.assertRaises(argparse.ArgumentTypeError):
                parse_coverage(value)


class ConfirmOrExitTest(unittest.TestCase):
    def accept(self, typed: str) -> bool:
        with patch('builtins.input', return_value=typed):
            try:
                confirm_or_exit('go?')
            except SystemExit:
                return False
        return True

    def test_an_empty_answer_accepts(self):
        self.assertTrue(self.accept(''))

    def test_yes_accepts(self):
        for typed in ('y', 'Y', 'yes', ' yes '):
            self.assertTrue(self.accept(typed), typed)

    def test_anything_else_stops(self):
        for typed in ('n', 'no', 'q', 'nope'):
            self.assertFalse(self.accept(typed), typed)


class PickBusinessUnitTest(unittest.TestCase):
    def setUp(self):
        self.first = {'name': 'Alpha', 'businessUnitId': 'a'}
        self.second = {'name': 'Beta', 'businessUnitId': 'b'}

    def pick(self, business_units, typed):
        with patch('builtins.input', return_value=typed):
            return pick_business_unit(business_units)

    def test_a_lone_business_unit_is_the_default(self):
        self.assertEqual(self.first, self.pick([self.first], ''))

    def test_the_first_is_the_default_when_there_are_several(self):
        self.assertEqual(self.first, self.pick([self.first, self.second], ''))

    def test_a_number_picks_that_business_unit(self):
        self.assertEqual(self.second, self.pick([self.first, self.second], '2'))

    def test_an_out_of_range_answer_is_rejected_then_retried(self):
        with patch('builtins.input', side_effect=['0', '7', 'beta', '2']):
            self.assertEqual(self.second, pick_business_unit([self.first, self.second]))


class PickAssessmentTest(unittest.TestCase):
    def setUp(self):
        self.first = {'assessmentName': 'first', 'assessmentId': 'a', 'assessmentTypeId': 'CIS 8.1'}
        self.second = {'assessmentName': 'second', 'assessmentId': 'b', 'assessmentTypeId': 'CIS 8.1'}

    def pick(self, existing, typed):
        with patch('builtins.input', return_value=typed):
            return pick_assessment(existing, 'CIS 8.1')

    def test_no_assessments_means_create_one(self):
        self.assertIsNone(self.pick([], ''))

    def test_a_lone_assessment_is_the_default(self):
        self.assertEqual(self.first, self.pick([self.first], ''))

    def test_the_first_is_the_default_when_there_are_several(self):
        self.assertEqual(self.first, self.pick([self.first, self.second], ''))

    def test_a_number_picks_that_assessment(self):
        self.assertEqual(self.second, self.pick([self.first, self.second], '2'))

    def test_the_last_option_creates_a_new_one(self):
        self.assertIsNone(self.pick([self.first, self.second], '3'))

    def test_an_out_of_range_answer_is_rejected_then_retried(self):
        with patch('builtins.input', side_effect=['9', 'nope', '2']):
            self.assertEqual(self.second, pick_assessment([self.first, self.second], 'CIS 8.1'))


class ApiTimestampTest(unittest.TestCase):
    def test_matches_the_layout_the_api_validates_against(self):
        stamp = api_timestamp(datetime(2026, 9, 2, 22, 30, 15, 123456, tzinfo=timezone.utc))
        self.assertEqual('2026-09-02T22:30:15.123456000Z', stamp)

    def test_always_nine_fractional_digits_even_on_a_whole_second(self):
        stamp = api_timestamp(datetime(2026, 1, 1, tzinfo=timezone.utc))
        self.assertEqual('2026-01-01T00:00:00.000000000Z', stamp)
        self.assertRegex(stamp, r'^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\.\d{9}Z$')


class ReadWorkbookTest(unittest.TestCase):
    def test_converts_each_status(self):
        with TemporaryDirectory() as directory:
            path = write_workbook(directory, [
                ['1.1', 1, 'Met', '5 (81-100%)', ''],
                ['1.2', 1, 'Partially Met', '4 (61-80%)', 'partial rollout'],
                ['1.3', 1, 'Not Met', '1 (0-20%)', ''],
                ['1.4', 1, 'Not Assessed', '1 (0-20%)', ''],
                ['1.5', 1, 'Not Applicable', None, ''],
            ])
            rows = read_workbook(path, 'CIS-Internal Controls')

        self.assertEqual(['1.1', '1.2', '1.3', '1.4', '1.5'], [r.number for r in rows])
        self.assertEqual([(SCORE, 100), (SCORE, 75), (SCORE, 0), (SCORE, 0), (SKIPPED, None)],
                         [(r.kind, r.score) for r in rows])
        self.assertEqual('partial rollout', rows[1].observations)

    def test_status_is_case_insensitive(self):
        with TemporaryDirectory() as directory:
            path = write_workbook(directory, [['1.1', 1, 'MET', '', '']])
            rows = read_workbook(path, 'CIS-Internal Controls')
        self.assertEqual((SCORE, 100), (rows[0].kind, rows[0].score))

    def test_stops_on_unrecognised_status(self):
        with TemporaryDirectory() as directory:
            path = write_workbook(directory, [
                ['1.1', 1, 'Met', '', ''],
                ['1.2', 1, 'Mostly There', '', ''],
            ])
            with self.assertRaises(SystemExit) as raised:
                read_workbook(path, 'CIS-Internal Controls')
        self.assertIn('1.2: Mostly There', str(raised.exception))

    def test_stops_on_blank_status(self):
        with TemporaryDirectory() as directory:
            path = write_workbook(directory, [['1.1', 1, None, '5 (81-100%)', '']])
            with self.assertRaises(SystemExit) as raised:
                read_workbook(path, 'CIS-Internal Controls')
        self.assertIn('(blank)', str(raised.exception))

    def test_stops_on_duplicate_safeguard(self):
        with TemporaryDirectory() as directory:
            path = write_workbook(directory, [
                ['1.1', 1, 'Met', '', ''],
                ['1.1', 1, 'Not Met', '', ''],
            ])
            with self.assertRaises(SystemExit) as raised:
                read_workbook(path, 'CIS-Internal Controls')
        self.assertIn('1.1', str(raised.exception))

    def test_stops_when_a_required_column_is_missing(self):
        with TemporaryDirectory() as directory:
            workbook = Workbook()
            workbook.active.title = 'CIS-Internal Controls'
            workbook.active.append(['CIS Safeguard', 'Observations'])
            workbook.active.append(['1.1', ''])
            path = str(Path(directory) / 'workbook.xlsx')
            workbook.save(path)
            with self.assertRaises(SystemExit) as raised:
                read_workbook(path, 'CIS-Internal Controls')
        self.assertIn('Assessment Status', str(raised.exception))

    def test_skips_rows_with_no_safeguard_number(self):
        with TemporaryDirectory() as directory:
            path = write_workbook(directory, [
                ['1.1', 1, 'Met', '', ''],
                [None, None, None, None, None],
            ])
            rows = read_workbook(path, 'CIS-Internal Controls')
        self.assertEqual(1, len(rows))

    def test_band_column_is_optional(self):
        with TemporaryDirectory() as directory:
            workbook = Workbook()
            workbook.active.title = 'Sheet1'
            workbook.active.append(['CIS Safeguard', 'Assessment Status'])
            workbook.active.append(['1.1', 'Met'])
            path = str(Path(directory) / 'workbook.xlsx')
            workbook.save(path)
            rows = read_workbook(path, 'Sheet1')
        self.assertEqual((SCORE, 100), (rows[0].kind, rows[0].score))


class BandDisagreesTest(unittest.TestCase):
    def test_matching_band_agrees(self):
        self.assertFalse(band_disagrees('5 (81-100%)', 'Met'))
        self.assertFalse(band_disagrees('4 (61-80%)', 'Partially Met'))
        self.assertFalse(band_disagrees('1 (0-20%)', 'Not Met'))

    def test_band_one_covers_both_not_met_and_not_assessed(self):
        self.assertFalse(band_disagrees('1 (0-20%)', 'Not Assessed'))

    def test_hand_edited_band_disagrees(self):
        self.assertTrue(band_disagrees('5 (81-100%)', 'Not Met'))
        self.assertTrue(band_disagrees('2 (21-40%)', 'Met'))

    def test_missing_or_unparseable_band_never_disagrees(self):
        self.assertFalse(band_disagrees('', 'Met'))
        self.assertFalse(band_disagrees('n/a', 'Met'))
        self.assertFalse(band_disagrees('3 (41-60%)', 'Not Applicable'))


class CoverageScoreTest(unittest.TestCase):
    def test_known_coverage_sends_a_score(self):
        self.assertEqual({'aspectId': 'ID.01-01-COVERAGE', 'score': 80},
                         coverage_score('ID.01-01', 80))

    def test_unknown_coverage_sends_unknown_only(self):
        entry = coverage_score('ID.01-01', UNKNOWN)
        self.assertEqual({'aspectId': 'ID.01-01-COVERAGE', 'unknown': True}, entry)
        self.assertNotIn('score', entry)


class BuildScoresTest(unittest.TestCase):
    def setUp(self):
        self.questions = {'1.1': 'ID.01-01', '1.2': 'PR.01-02', '1.3': 'DE.01-03', '1.4': 'RS.01-04'}

    def test_pairs_coverage_with_every_process_score(self):
        rows = [Row('1.1', 'Met', SCORE, 100, ''), Row('1.2', 'Not Met', SCORE, 0, '')]
        self.assertEqual([
            {'aspectId': 'ID.01-01-PROCESS', 'score': 100},
            {'aspectId': 'ID.01-01-COVERAGE', 'score': 80},
            {'aspectId': 'PR.01-02-PROCESS', 'score': 0},
            {'aspectId': 'PR.01-02-COVERAGE', 'score': 80},
        ], build_scores(rows, self.questions, 80))

    def test_not_assessed_scores_zero_like_any_other_scored_row(self):
        rows = [Row('1.4', 'Not Assessed', SCORE, 0, '')]
        self.assertEqual([
            {'aspectId': 'RS.01-04-PROCESS', 'score': 0},
            {'aspectId': 'RS.01-04-COVERAGE', 'score': 80},
        ], build_scores(rows, self.questions, 80))

    def test_unanswered_process_still_gets_coverage(self):
        rows = [Row('1.4', 'Not Assessed', UNKNOWN, None, '')]
        self.assertEqual([
            {'aspectId': 'RS.01-04-PROCESS', 'unknown': True},
            {'aspectId': 'RS.01-04-COVERAGE', 'score': 80},
        ], build_scores(rows, self.questions, 80))

    def test_not_applicable_is_left_out_entirely(self):
        rows = [Row('1.3', 'Not Applicable', SKIPPED, None, '')]
        self.assertEqual([], build_scores(rows, self.questions, 80))

    def test_a_skipped_row_does_not_stop_the_others(self):
        rows = [Row('1.3', 'Not Applicable', SKIPPED, None, ''), Row('1.1', 'Met', SCORE, 100, '')]
        self.assertEqual([
            {'aspectId': 'ID.01-01-PROCESS', 'score': 100},
            {'aspectId': 'ID.01-01-COVERAGE', 'score': 80},
        ], build_scores(rows, self.questions, 80))


class BuildNotesTest(unittest.TestCase):
    def test_only_rows_with_observations_become_notes(self):
        rows = [Row('1.1', 'Met', SCORE, 100, 'reviewed by internal audit'),
                Row('1.2', 'Met', SCORE, 100, '')]
        self.assertEqual([{'aspectId': 'ID.01-01-PROCESS', 'value': 'reviewed by internal audit'}],
                         build_notes(rows, {'1.1': 'ID.01-01', '1.2': 'PR.01-02'}))


class BuildSafeguardMapTest(unittest.TestCase):
    def test_maps_safeguard_number_parsed_from_question_text(self):
        content = {'nodes': {
            'ID.01-01': {'text': 'CIS Safeguard 1.1 Establish and Maintain', 'process': {}},
            'PR.02-05': {'text': 'CIS Safeguard 2.5 Allowlist Authorized Software', 'process': {}},
            'ID.01': {'text': 'Inventory and Control of Enterprise Assets'},
        }}
        self.assertEqual({'1.1': 'ID.01-01', '2.5': 'PR.02-05'}, build_safeguard_map(content))

    def test_empty_content_maps_nothing(self):
        self.assertEqual({}, build_safeguard_map({}))

    def test_maps_the_real_content_spec_shape(self):
        """The spec keys question nodes by id and carries the safeguard number in the text.

        Mirrors the CIS 8.1 content the api serves, em dash and all, so the pattern is not
        quietly narrowed to a punctuation style the real content does not use.
        """
        content = {
            'descriptor': {'assessment-key': 'CIS 8.1'},
            'questions': {'GV': {'question-count': 1}},
            'colors': {},
            'nodes': {
                'GV': {'name': 'Govern', 'text': 'Govern'},
                'GV.03-01': {
                    'text': 'Safeguard 3.1 — Establish and Maintain a Data Management Process',
                    'process': [{'text': 'Fully implemented', 'value': 100}],
                    'coverage': 'What share of the organisation does this cover?',
                },
            },
        }
        self.assertEqual({'3.1': 'GV.03-01'}, build_safeguard_map(content))


if __name__ == '__main__':
    unittest.main()
